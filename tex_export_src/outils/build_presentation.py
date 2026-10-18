"""Construit la présentation de soutenance du rapport de stage (skill latex-pdf-soutenance).

Usage : python outils/build_presentation.py <dossier tex_export_src> <sortie.pptx>

- Les informations (noms, entreprise, chantier…) sont lues dans config/infos.tex :
  une valeur \\AC{x} s'affiche « [À COMPLÉTER : x] » en orange, comme dans le PDF.
- Les schémas TikZ de figures/ sont compilés en PNG (outils/_build/) et insérés tels quels.
- Les images factices de img/ sont réutilisées : remplacer la photo dans img/ suffit.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = Path(sys.argv[1]).resolve()
OUT = Path(sys.argv[2]).resolve()
BUILD = SRC / "outils" / "_build"

VERT = RGBColor(0x2E, 0x7D, 0x32)
NUIT = RGBColor(0x1B, 0x4D, 0x1F)
MOYEN = RGBColor(0x66, 0xBB, 0x6A)
CLAIR = RGBColor(0xEE, 0xF6, 0xEC)
ORANGE = RGBColor(0xE6, 0x51, 0x00)
GRIS = RGBColor(0x5B, 0x67, 0x70)
TEXTE = RGBColor(0x22, 0x2B, 0x33)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
SECTIONS = ["Introduction", "Entreprise", "Stage", "Plomberie", "Critiques", "Conclusion"]


# ------------------------------------------------------------------ données
def lire_infos() -> dict[str, str]:
    """Lit les macros \\newcommand{\\InfoX}{valeur} de config/infos.tex."""
    txt = (SRC / "config" / "infos.tex").read_text(encoding="utf-8")
    infos = {}
    for m in re.finditer(r"\\newcommand\{\\(Info\w+)\}\{(.*)\}\s*(?:%.*)?$", txt, re.M):
        val = m.group(2)
        val = re.sub(r"\\AC\{(.*)\}", r"[À COMPLÉTER : \1]", val)
        val = val.replace(r"\textsuperscript{e}", "e").replace("{", "").replace("}", "").replace("\\", "")
        infos[m.group(1)] = val
    return infos


I = lire_infos()


def compiler_figures() -> dict[str, Path]:
    """Compile chaque figures/*.tex en PNG via une classe standalone."""
    BUILD.mkdir(parents=True, exist_ok=True)
    out = {}
    for fig in sorted((SRC / "figures").glob("*.tex")):
        doc = BUILD / f"{fig.stem}_standalone.tex"
        doc.write_text(rf"""\documentclass[border=6pt]{{standalone}}
\usepackage[T1]{{fontenc}}\usepackage[utf8]{{inputenc}}\usepackage[french]{{babel}}
\usepackage[scaled=0.95]{{helvet}}\renewcommand{{\familydefault}}{{\sfdefault}}
\usepackage{{xcolor,tikz,amsmath,amssymb}}\usepackage[locale=FR]{{siunitx}}
\usetikzlibrary{{arrows.meta,positioning,shapes.geometric,calc,fit,shadows,decorations.pathmorphing}}
\definecolor{{vertnuit}}{{HTML}}{{1B4D1F}}\definecolor{{vertfonce}}{{HTML}}{{2E7D32}}
\definecolor{{vertmoyen}}{{HTML}}{{66BB6A}}\definecolor{{vertclair}}{{HTML}}{{DCEFD9}}
\definecolor{{bandeau}}{{HTML}}{{DDE1F0}}\definecolor{{acompleter}}{{HTML}}{{E65100}}
\newcommand{{\AC}}[1]{{\textcolor{{acompleter}}{{[À COMPLÉTER : #1]}}}}
\begin{{document}}\input{{figures/{fig.stem}}}\end{{document}}
""", encoding="utf-8")
        subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error",
                        f"-output-directory={BUILD}", str(doc)], cwd=SRC,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
        pdf = fitz.open(BUILD / f"{fig.stem}_standalone.pdf")
        png = BUILD / f"{fig.stem}.png"
        pdf[0].get_pixmap(dpi=220).save(png)
        out[fig.stem] = png
    return out


FIG = compiler_figures()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def style_run(run, size, bold=False, color=TEXTE, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_runs(p, text, size, bold, color, italic=False):
    """Écrit le texte en colorant les segments [À COMPLÉTER : …] en orange."""
    for part in re.split(r"(\[À COMPLÉTER : [^\]]*\])", text):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        style_run(r, size, bold, ORANGE if part.startswith("[À COMPLÉTER") else color, italic)


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=TEXTE, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        add_runs(p, line, size, bold, color, italic)
    return box


def bullets(slide, x, y, w, h, items, size=17, gap=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "◆  "; style_run(r, size - 4, True, VERT)
        if isinstance(item, tuple):
            add_runs(p, item[0], size, True, NUIT)
            add_runs(p, item[1], size, False, TEXTE)
        else:
            add_runs(p, item, size, False, TEXTE)
    return box


def rect(slide, x, y, w, h, fill=CLAIR, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.5)
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, title, items, size=15, head=VERT):
    rect(slide, x, y, w, h, CLAIR)
    rect(slide, x, y, w, 0.09, head)
    textbox(slide, x + 0.15, y + 0.17, w - 0.3, 0.45, title, 17, True, NUIT if head == VERT else head)
    bullets(slide, x + 0.12, y + 0.7, w - 0.25, h - 0.8, items, size, gap=4)


def metric(slide, x, y, w, value, label, color=VERT):
    rect(slide, x, y, w, 1.3, BLANC, line=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(slide, x, y + 0.08, w, 0.65, value, 28, True, color, PP_ALIGN.CENTER)
    textbox(slide, x + 0.1, y + 0.75, w - 0.2, 0.5, label, 13, False, GRIS, PP_ALIGN.CENTER)


def image(slide, path, x, y, w, h):
    """Insère une image ajustée dans la boîte (x, y, w, h), sans déformation."""
    from PIL import Image
    path = Path(path)
    if path.suffix == ".png" and not path.exists() and path.with_suffix(".jpg").exists():
        path = path.with_suffix(".jpg")  # photo réelle importée à la place de l'image factice
    with Image.open(path) as im:
        iw, ih = im.size
    r = min(w / iw, h / ih)
    nw, nh = iw * r, ih * r
    return slide.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2),
                                    Inches(nw), Inches(nh))


def table(slide, x, y, w, rows, col_w, size=14, row_h=0.5):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.fill.solid()
            cell.fill.fore_color.rgb = VERT if r == 0 else (CLAIR if r % 2 else BLANC)
            tf = cell.text_frame; tf.clear(); tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 or nc == 2 else PP_ALIGN.CENTER
            add_runs(p, str(rows[r][c]), size, r == 0 or c == 0, BLANC if r == 0 else TEXTE)


def banniere(slide, x, y, w, h, text, size=26):
    b = rect(slide, x, y, w, h, MOYEN, line=VERT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, line, size if i == 0 else size - 6, i == 0, BLANC)
    return b


def nouvelle(title, section=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 0.32, RGBColor(0xDD, 0xE1, 0xF0))
    textbox(s, 0.4, 0.02, 9, 0.3, f"{I['InfoTheme']} — cas de {I['InfoEntrepriseSigle']}", 10, True, NUIT, italic=True)
    if section:
        nav = s.shapes.add_textbox(Inches(7.0), Inches(0.02), Inches(6.0), Inches(0.3))
        p = nav.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
        for i, name in enumerate(SECTIONS):
            r = p.add_run(); r.text = name
            style_run(r, 10, name == section, ORANGE if name == section else GRIS)
            if i < len(SECTIONS) - 1:
                r = p.add_run(); r.text = "  ·  "; style_run(r, 10, False, GRIS)
    textbox(s, 0.5, 0.5, 12.3, 0.8, title, 28, True, NUIT, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0.55, 1.28, 1.3, 0.07, VERT)
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(6.98), Inches(12.85), Inches(6.98)).line.color.rgb = RGBColor(0xC8, 0xD8, 0xC8)
    textbox(s, 0.5, 7.02, 9, 0.35, f"Rédigé et présenté par {I['InfoEtudiantCourt']}  —  {I['InfoSpecialite']}", 11, True, GRIS, italic=True)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def numeroter():
    total = len(prs.slides)
    for i, s in enumerate(prs.slides, 1):
        if i in (1, total):
            continue
        r = rect(s, 12.0, 7.02, 0.85, 0.32, VERT)
        tf = r.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = f"{i} / {total}"; style_run(run, 11, True, BLANC)


# ------------------------------------------------------------------ 1. Couverture
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, BLANC)
rect(s, 0.25, 0.25, 12.833, 7.0, BLANC, line=VERT)
textbox(s, 0.6, 0.45, 5.5, 0.9, "RÉPUBLIQUE DU CAMEROUN\nMinistère de l'Emploi et de la Formation professionnelle", 11, True, GRIS)
textbox(s, 7.2, 0.45, 5.5, 0.9, "REPUBLIC OF CAMEROON\nMinistry of Employment and Vocational Training", 11, True, GRIS, PP_ALIGN.RIGHT)
image(s, SRC / "img" / "logo_cfp.png", 0.7, 1.45, 1.4, 1.4)
image(s, SRC / "img" / "logo_entreprise.png", 11.2, 1.45, 1.4, 1.4)
textbox(s, 2.3, 1.5, 8.7, 0.5, "EXAMENS NATIONAUX DE FIN DE CYCLE DE FORMATION PROFESSIONNELLE", 14, True, NUIT, PP_ALIGN.CENTER)
textbox(s, 2.3, 2.0, 8.7, 0.5, "Soutenance du rapport de stage", 18, False, GRIS, PP_ALIGN.CENTER)
banniere(s, 1.4, 2.9, 10.5, 1.55, f"Thème : {I['InfoTheme'].upper()}\n{I['InfoSousTheme']}", 30)
textbox(s, 0.8, 4.65, 11.7, 0.5, f"Spécialité : {I['InfoSpecialite']}  —  stage effectué à {I['InfoEntreprise']}", 15, True, NUIT, PP_ALIGN.CENTER)
textbox(s, 0.8, 5.25, 5.6, 1.2, f"Présenté par\n{I['InfoEtudiant']}", 15, True, TEXTE, PP_ALIGN.CENTER)
textbox(s, 6.9, 5.25, 5.6, 1.2, f"Encadrement\n{I['InfoTuteurAcad']} (académique)\n{I['InfoTuteurPro']} (professionnel)", 13, False, TEXTE, PP_ALIGN.CENTER)
textbox(s, 0.8, 6.55, 11.7, 0.4, f"Session {I['InfoSession']}", 13, True, VERT, PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = ("Bonjour. Je vous présente mon rapport de stage de fin de formation en froid et "
    "climatisation, dont le thème est la plomberie sanitaire : la distribution d'eau froide et d'eau chaude dans un immeuble à étages.")

# ------------------------------------------------------------------ 2. Plan
s = nouvelle("Plan de la présentation", notes="La présentation suit le plan du rapport.")
plan = [("Introduction", "Contexte et choix du stage"), ("L'entreprise d'accueil", "Activités, organisation, environnement"),
        ("Déroulement du stage", "Poste, planning, résultats, compétences"),
        ("Plomberie sanitaire", "Généralités, dimensionnement, mise en œuvre, liens avec le froid"),
        ("Critiques et suggestions", "Points forts et axes de progrès"), ("Conclusion", "Bilan et projet professionnel")]
for i, (t, d) in enumerate(plan):
    y = 1.6 + i * 0.86
    b = rect(s, 0.9, y, 0.7, 0.7, VERT if i % 2 == 0 else MOYEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f"{i + 1}"; style_run(r, 20, True, BLANC)
    textbox(s, 1.85, y - 0.02, 10, 0.4, t, 20, True, NUIT)
    textbox(s, 1.85, y + 0.36, 10, 0.35, d, 14, False, GRIS)

# ------------------------------------------------------------------ 3. Introduction
s = nouvelle("Contexte et choix du stage", "Introduction", notes=(
    "La construction urbaine se développe au Cameroun et chaque bâtiment a besoin d'installations techniques fiables. "
    "J'ai choisi une entreprise multi-lots pour voir comment froid, plomberie et électricité s'articulent sur un même chantier."))
card(s, 0.5, 1.6, 4.0, 5.1, "Contexte", ["Forte croissance de la construction urbaine",
     "Lots techniques : confort, hygiène, sécurité", "Demande en techniciens qualifiés"])
card(s, 4.67, 1.6, 4.0, 5.1, "Pourquoi ce stage ?", ["Entreprise multi-lots : froid, plomberie, électricité",
     "Grands chantiers d'immeubles", "Gestes communs : tubes, assemblages, essais"])
card(s, 8.83, 1.6, 4.0, 5.1, "Le thème", [("Plomberie sanitaire", " : eau froide et eau chaude"),
     f"Chantier {I['InfoChantier']}", "Lien direct avec le froid : isolation, condensats, ECS"], head=ORANGE)

# ------------------------------------------------------------------ 4. Entreprise
s = nouvelle("L'entreprise d'accueil", "Entreprise", notes=(
    "L'entreprise est une SARL spécialisée dans les lots techniques du bâtiment. Elle va de l'étude à la maintenance."))
textbox(s, 0.5, 1.5, 12.3, 0.8, f"{I['InfoEntreprise']} ({I['InfoEntrepriseSigle']}) — SARL créée en {I['InfoEntrepriseCreation']}, siège à {I['InfoEntrepriseVille']}", 16, True, NUIT)
doms = [("Climatisation", "split, cassette, DRV, ventilation"), ("Froid", "chambres froides, maintenance"),
        ("Plomberie", "EF, ECS, évacuations, surpresseurs"), ("Électricité", "courants forts et faibles"),
        ("Incendie", "détection, extinction, RIA")]
for i, (t, d) in enumerate(doms):
    x = 0.5 + i * 2.5
    rect(s, x, 2.45, 2.3, 1.6, CLAIR); rect(s, x, 2.45, 2.3, 0.08, VERT if i != 2 else ORANGE)
    textbox(s, x + 0.1, 2.62, 2.1, 0.45, t, 17, True, NUIT, PP_ALIGN.CENTER)
    textbox(s, x + 0.1, 3.1, 2.1, 0.9, d, 13, False, TEXTE, PP_ALIGN.CENTER)
card(s, 0.5, 4.3, 6.05, 2.45, "Forces", ["Offre multi-lots, interlocuteur unique", "Personnel expérimenté, grands chantiers", "Activité de maintenance"], 14)
card(s, 6.75, 4.3, 6.05, 2.45, "Faiblesses et menaces", ["Logistique et stocks perfectibles", "Sécurité peu formalisée", "Concurrence informelle, coût des importations"], 14, head=ORANGE)

# ------------------------------------------------------------------ 5. Organisation
s = nouvelle("Organisation et poste occupé", "Entreprise", notes=(
    "L'entreprise est organisée en divisions. J'étais aide-technicien plombier, rattaché à l'équipe plomberie."))
image(s, FIG["organigramme"], 0.4, 1.5, 8.4, 5.3)
card(s, 9.0, 1.6, 3.85, 5.1, "Mon poste", [("Aide-technicien plombier", ""), "Division plomberie, sous le chef d'équipe",
     f"Tuteur : {I['InfoTuteurPro']}", f"Durée : {I['InfoDuree']}"], 14)

# ------------------------------------------------------------------ 6. Déroulement
s = nouvelle("Déroulement du stage", "Stage", notes=(
    "Le stage s'est déroulé en phases : préparation, pose des canalisations, essais, isolation, réducteurs de pression, "
    "robinetterie, appui à la climatisation, puis mise en eau."))
image(s, FIG["gantt"], 0.4, 1.5, 8.6, 4.4)
card(s, 9.2, 1.6, 3.65, 5.1, "Résultats", ["Réseaux EF/ECS posés et testés", "Réducteurs et manomètres montés", "Canalisations isolées", "Essais d'étanchéité conformes"], 14)
textbox(s, 0.5, 6.1, 8.4, 0.7, "Organisation étage par étage : traçage → supports → assemblage → raccordement → essai → autocontrôle", 14, True, NUIT)

# ------------------------------------------------------------------ 7. Compétences
s = nouvelle("Difficultés et compétences acquises", "Stage", notes=(
    "Chaque difficulté a trouvé une solution. J'ai acquis des compétences techniques et personnelles."))
table(s, 0.5, 1.6, 7.3, [["Difficulté", "Solution"],
      ["Polyfusion irrégulière au début", "Entraînement, respect des temps de chauffe"],
      ["Lecture des plans multi-niveaux", "Accompagnement, repérage des cotes"],
      ["Ruptures de raccords", "Anticipation, réorganisation des zones"],
      ["Hauteur et faible éclairage", "Harnais, projecteurs"]], [3.4, 3.9], 13, 0.85)
card(s, 8.05, 1.6, 4.8, 5.1, "Compétences", ["Lire un plan de plomberie", "Assembler le PPR par polyfusion",
     "Monter et régler un réducteur", "Réaliser un essai d'étanchéité", "Rigueur, ponctualité, travail d'équipe"], 14)

# ------------------------------------------------------------------ 8. Généralités
s = nouvelle("Plomberie sanitaire : principe d'installation", "Plomberie", notes=(
    "Une installation comprend le branchement, éventuellement une bâche et un surpresseur, les colonnes d'eau froide et "
    "d'eau chaude, les réducteurs de pression aux étages, les appareils et les évacuations."))
image(s, FIG["schema_principe"], 0.4, 1.5, 8.6, 5.3)
card(s, 9.2, 1.6, 3.65, 5.1, "Matériaux", [("PPR", " : EF et ECS, polyfusion"), ("PER", " : distribution"), ("PE", " : branchement enterré"), ("PVC", " : évacuations")], 14)

# ------------------------------------------------------------------ 9. Dimensionnement
s = nouvelle("Étude de cas : dimensionnement de la colonne", "Plomberie", notes=(
    "Exemple d'un immeuble R+10, 40 chambres avec lavabo, douche et WC. Avec le coefficient de simultanéité du DTU 60.11 "
    "majoré pour l'hôtellerie, le débit probable de la colonne est de 1,91 litre par seconde. Avec une vitesse maximale de "
    "1,5 mètre par seconde, on retient un tube PPR de 63 millimètres."))
textbox(s, 0.5, 1.5, 7.4, 0.9, "Hypothèses : R+10, 40 chambres (lavabo 0,20 + douche 0,20 + WC 0,12 L/s), usage hôtelier Cm = 1,25", 15, False, TEXTE)
textbox(s, 0.5, 2.35, 7.4, 0.6, "y = 0,8 / √(N − 1) × Cm        Qp = y × ΣQb        D = √(4 Qp / π V)", 17, True, NUIT)
table(s, 0.5, 3.15, 7.4, [["Tronçon", "Qp (L/s)", "Vmax (m/s)", "PPR retenu", "V (m/s)"],
      ["Étage (N = 12)", "0,63", "1,0", "50 × 8,3", "0,72"],
      ["Colonne (N = 120)", "1,91", "1,5", "63 × 10,5", "1,38"]], [2.2, 1.2, 1.3, 1.5, 1.2], 14, 0.6)
textbox(s, 0.5, 5.2, 7.4, 1.0, "Pertes de charge (Flamant) ≈ 3,4 m + hauteur 30 m + 1 bar résiduel ⇒ ≈ 4,3 bar en pied de colonne", 14, False, TEXTE)
metric(s, 8.3, 1.6, 2.15, "1,91 L/s", "débit probable colonne")
metric(s, 10.7, 1.6, 2.15, "PPR 63", "diamètre de colonne")
metric(s, 8.3, 3.15, 2.15, "1,38 m/s", "vitesse ≤ 1,5 m/s", ORANGE)
metric(s, 10.7, 3.15, 2.15, "4,3 bar", "pression en pied")
textbox(s, 8.3, 4.75, 4.6, 1.5, "Source : NF DTU 60.11 (débits de base, simultanéité, vitesses)", 12, False, GRIS, italic=True)

# ------------------------------------------------------------------ 10. Réducteurs
s = nouvelle("Pourquoi des réducteurs de pression ?", "Plomberie", notes=(
    "La pression décroît d'environ 0,3 bar par étage. Du rez-de-chaussée au quatrième étage, elle dépasse 3 bar : "
    "on y installe un réducteur de pression réglé à 3 bar, avec un manomètre de contrôle."))
image(s, FIG["colonne_montante"], 0.4, 1.5, 7.2, 5.3)
card(s, 7.9, 1.6, 4.95, 5.1, "Justification", ["−0,3 bar par étage", "RDC à R+4 : P > 3 bar", "Réducteur réglé à 3 bar + manomètre",
     "Rapport de réduction < 3:1 : pas de cavitation", "Protège la robinetterie, limite la consommation"], 14, head=ORANGE)

# ------------------------------------------------------------------ 11. Mise en œuvre
s = nouvelle("Mise en œuvre sur le chantier", "Plomberie", notes=(
    "Quatre opérations clés : la polyfusion du PPR à 260 degrés, le montage des réducteurs, l'isolation et l'essai d'étanchéité."))
ops = [("Polyfusion PPR", "fig_polyfuseur", ["Matrices à 260 °C", "Coupe d'équerre, repère", "Chauffe, emboîtement sans rotation"]),
       ("Réducteur + manomètre", "fig_reducteur_montage", ["Sens de la flèche", "Entre deux vannes", "Réglage à 3 bar"]),
       ("Isolation", "fig_isolation_pose", ["Mousse à cellules fermées", "Collage des fentes", "Anti-condensation"]),
       ("Essai d'étanchéité", "fig_essai_etancheite", ["Purge, mise en pression", "Durée prescrite", "Fiche d'autocontrôle"])]
for i, (t, img, items) in enumerate(ops):
    x = 0.5 + i * 3.1
    rect(s, x, 1.55, 2.95, 5.2, CLAIR); rect(s, x, 1.55, 2.95, 0.08, VERT)
    textbox(s, x + 0.1, 1.7, 2.75, 0.45, t, 16, True, NUIT, PP_ALIGN.CENTER)
    image(s, SRC / "img" / f"{img}.png", x + 0.15, 2.2, 2.65, 1.65)
    bullets(s, x + 0.1, 4.0, 2.8, 2.6, items, 13, 3)

# ------------------------------------------------------------------ 12. Liens froid
s = nouvelle("Liens avec le froid et la climatisation", "Plomberie", notes=(
    "Le thème rejoint ma spécialité : l'isolation évite la condensation comme sur les lignes frigorifiques, les condensats "
    "des splits rejoignent le réseau d'eaux usées, et une pompe à chaleur peut produire l'eau chaude."))
image(s, FIG["cycle_frigorifique"], 0.4, 1.5, 7.2, 5.3)
card(s, 7.9, 1.6, 4.95, 5.1, "Interfaces plomberie / froid", [("Isolation", " : surface au-dessus du point de rosée"),
     ("Condensats", " : siphon, pente ≥ 1 %, réseau EU"), ("ECS", " : pompe à chaleur, récupération de chaleur"),
     ("Appui", " à l'équipe climatisation (S15)")], 14)

# ------------------------------------------------------------------ 13. Critiques
s = nouvelle("Critiques et suggestions", "Critiques", notes=(
    "L'entreprise a un bon encadrement technique. Les axes de progrès sont la logistique, la sécurité et la traçabilité."))
table(s, 0.5, 1.6, 12.35, [["Constat", "Suggestion"],
      ["Ruptures de raccords et consommables", "Liste de matériel par zone, stock tampon"],
      ["Disparitions de matériel", "Local fermé par niveau, registre des sorties"],
      ["EPI pas toujours portés, travaux en hauteur", "Référent HSE, quart d'heure sécurité, garde-corps"],
      ["Éclairage insuffisant", "Projecteurs et éclairage provisoire"],
      ["Autocontrôles irréguliers", "Fiches signées, photos des essais"]], [5.6, 6.75], 14, 0.72)
textbox(s, 0.5, 6.15, 12.3, 0.6, "Points forts : encadrement technique, organisation par zones, essais systématiques, coordination entre lots", 14, True, VERT)

# ------------------------------------------------------------------ 14. Conclusion
s = nouvelle("Conclusion et perspectives", "Conclusion", notes=(
    "Le stage a été très enrichissant. Il a conforté mon projet : devenir technicien en froid et climatisation, capable "
    "d'intervenir aussi sur les réseaux de fluides du bâtiment."))
card(s, 0.5, 1.6, 4.0, 5.1, "Bilan", ["Gestes professionnels acquis", "Théorie (DTU 60.11) appliquée au chantier", "Rigueur et travail d'équipe"])
card(s, 4.67, 1.6, 4.0, 5.1, "Limites et prolongements", ["Évacuations non traitées", "Hypothèses de calcul simplifiées", "Étude énergétique de l'ECS"], head=ORANGE)
card(s, 8.83, 1.6, 4.0, 5.1, "Projet professionnel", ["Technicien froid et climatisation", "Compétences fluides du bâtiment", "Insertion en entreprise multi-lots"])

# ------------------------------------------------------------------ 15. Merci
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, VERT)
rect(s, 5.67, 4.05, 2.0, 0.06, BLANC)
textbox(s, 0.6, 2.3, 12.1, 1.2, "Merci pour votre attention", 44, True, BLANC, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
textbox(s, 0.6, 3.35, 12.1, 0.6, "Questions et échanges", 24, False, CLAIR, PP_ALIGN.CENTER)
rect(s, 3.2, 4.35, 6.9, 0.6, BLANC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 3.2, 4.38, 6.9, 0.55, I["InfoEtudiant"], 18, True, NUIT, PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = "Je vous remercie pour votre attention et reste disponible pour vos questions."

numeroter()
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT, len(prs.slides), "diapositives")
